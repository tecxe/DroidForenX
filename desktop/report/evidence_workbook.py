"""XLSX verification workbook for a DroidForenX case.

Companion to the PDF Evidence Report. The workbook holds the same evidence data
in a form an examiner can sort, filter and re-check: a case summary, an
independent re-hash of every artifact, the per-source status and verification
tier, the artifact log, and the chain of custody.

Like the PDF, this module reads no AI output and contains none.

Building the workbook is a strictly read-only operation: it never writes to the
case tree and never appends a custody event. Integrity is recomputed with
``compute_integrity`` (a pure inspection) rather than ``verify`` (which records
the run), so producing a report can never alter evidence.

Usage:
    python -m desktop.report.evidence_workbook CASE-ID
    python -m desktop.report.evidence_workbook CASE-ID -o out.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from desktop.core import evidence_store  # noqa: E402
from desktop.core.evidence_store import AcquisitionStatus, VerificationTier  # noqa: E402
from desktop.report.evidence_report import _latest_status_per_loader, _load  # noqa: E402

# --------------------------------------------------------------------------
# styling
# --------------------------------------------------------------------------

INK = "FF1B1F27"
HEADER_BG = "FF1F2733"
ZEBRA_BG = "FFF4F6F8"
OK_BG = "FFDFF3E3"
WARN_BG = "FFFDF0D5"
BAD_BG = "FFFBDCDC"
MONO = "Consolas"

_HDR_FONT = Font(bold=True, color="FFFFFFFF", size=10)
_BODY = Font(size=10)
_BOLD = Font(size=10, bold=True)
_MONO = Font(name=MONO, size=9)
_TITLE = Font(bold=True, size=14, color=INK)
_SUB = Font(italic=True, size=9, color="FF5A6472")

_THIN = Side(style="thin", color="FFD3D9E0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(vertical="top", wrap_text=True)
_TOP = Alignment(vertical="top")

#: Semantic status colouring, matching the PDF report.
STATUS_FILLS = {
    AcquisitionStatus.ACQUIRED: OK_BG,
    AcquisitionStatus.PARTIAL: WARN_BG,
    AcquisitionStatus.RECOVERED: OK_BG,
    AcquisitionStatus.DELETED_RECOVERABLE: OK_BG,
    AcquisitionStatus.DELETED_UNRECOVERABLE: WARN_BG,
    AcquisitionStatus.NOT_ACCESSIBLE: BAD_BG,
    AcquisitionStatus.NOT_PRESENT: BAD_BG,
    AcquisitionStatus.FAILED: BAD_BG,
}

VERDICT_FILLS = {
    "VERIFIED": OK_BG,
    "INTEGRITY_FAILURE": BAD_BG,
    "NO_MANIFEST": WARN_BG,
}

TIER_FILLS = {
    VerificationTier.T1_DEVICE_VERIFIED: OK_BG,
    VerificationTier.T2_FIXTURE_VERIFIED: WARN_BG,
    VerificationTier.T3_UNTESTED: BAD_BG,
}

#: Per-row re-hash result colouring. Absent entries fall through to zebra.
RESULT_FILLS = {
    "OK": OK_BG,
    "MISMATCH": BAD_BG,
    "MISSING": BAD_BG,
}


def _header(ws, labels, widths):
    """Write a styled header row and apply column widths."""
    for idx, (label, width) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.font = _HDR_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 26


def _finish(ws, widths, first_data_row=2):
    ws.freeze_panes = ws.cell(row=first_data_row, column=1)
    if ws.max_row >= first_data_row and ws.max_column:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, first_data_row)}"
        )


def _write_row(ws, row, values, *, mono_cols=(), fill=None):
    for idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=idx, value=value)
        cell.font = _MONO if idx in mono_cols else _BODY
        cell.alignment = _WRAP if isinstance(value, str) and len(str(value)) > 48 else _TOP
        cell.border = _BORDER
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        elif row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=ZEBRA_BG)


# --------------------------------------------------------------------------
# sheets
# --------------------------------------------------------------------------

def _sheet_summary(wb, data, integrity, case_id):
    ws = wb.active
    ws.title = "Case Summary"

    meta = data["meta"]
    ws["A1"] = "DroidForenX — Case Verification Workbook"
    ws["A1"].font = _TITLE
    ws["A2"] = "Evidence only. Contains no AI output. Produced read-only; no evidence was altered."
    ws["A2"].font = _SUB
    ws["A3"] = ""

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 78

    row = 5
    ws.cell(row=row, column=1, value="Field").font = _HDR_FONT
    ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=HEADER_BG)
    ws.cell(row=row, column=2, value="Value").font = _HDR_FONT
    ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=HEADER_BG)
    for col in ("A", "B"):
        ws.cell(row=row, column={"A": 1, "B": 2}[col]).border = _BORDER

    statuses = data["statuses"]
    acquired = [s for s in statuses if AcquisitionStatus.is_success(s.get("status", ""))]
    unreachable = [s for s in statuses
                   if s.get("status") in AcquisitionStatus.UNSUCCESSFUL]
    verified_loaders = {v["loader"] for v in data["verifications"]
                        if v["tier"] == VerificationTier.T1_DEVICE_VERIFIED}

    rows = [
        ("Case ID", meta.get("case_id") or case_id),
        ("Operator", meta.get("operator") or "(not recorded)"),
        ("Device serial", meta.get("device_serial") or "(not recorded)"),
        ("Device model", meta.get("device_model") or "(not recorded)"),
        ("Case opened", meta.get("opened_at") or "(not recorded)"),
        ("Tool", f"{meta.get('tool', 'DroidForenX')} {meta.get('tool_version', '')}".strip()),
        ("Workbook generated", evidence_store._utc_now()),
        ("", ""),
        ("INTEGRITY", integrity["verdict"]),
        ("Artifacts in manifest", integrity["checked"]),
        ("Re-hash matched", integrity["ok"]),
        ("Hash mismatches", len(integrity["mismatched"])),
        ("Missing artifacts", len(integrity["missing"])),
        ("", ""),
        ("Sources recorded", len(statuses)),
        ("Sources reporting acquisition", len(acquired)),
        ("Sources not obtained", len(unreachable)),
        ("Loaders device-verified (T1)", len(verified_loaders)),
        ("Artifacts in artifact log", len(data["artifacts"])),
        ("Custody events", len(data["custody"])),
    ]

    for label, value in rows:
        ws.cell(row=row, column=1, value=label).font = _BOLD
        ws.cell(row=row, column=1).alignment = _TOP
        ws.cell(row=row, column=1).border = _BORDER
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = _BODY
        cell.alignment = _TOP
        cell.border = _BORDER
        if label == "INTEGRITY":
            cell.font = _BOLD
            cell.fill = PatternFill("solid", fgColor=VERDICT_FILLS.get(value, WARN_BG))
        row += 1

    return ws


def _sheet_integrity(wb, data, integrity, case_id):
    ws = wb.create_sheet("Integrity")
    _header(ws, ["Artifact path", "Manifest SHA-256", "Re-computed SHA-256", "Result"],
            [52, 68, 68, 18])

    if not data["hashes"]:
        note = ws.cell(
            row=2, column=1,
            value="NO_MANIFEST - no artifact has been acquired for this case yet. "
                  "Integrity cannot be asserted until an acquisition produces a "
                  "SHA256SUMS manifest.")
        note.font = _BOLD
        note.alignment = _WRAP
        note.fill = PatternFill("solid", fgColor=WARN_BG)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
        return ws

    row = 2
    actual_hashes = integrity.get("actual") or {}
    for digest, relative in data["hashes"]:
        actual = actual_hashes.get(relative)
        if relative in integrity["missing"]:
            # Nothing on disk to re-hash; leave the column empty rather than
            # echoing the manifest digest, which would read as a false match.
            result, recomputed = "MISSING", ""
        elif relative in integrity["mismatched"]:
            result, recomputed = "MISMATCH", (actual or "")
        else:
            result, recomputed = "OK", (actual or digest)
        # Colour by this row's own result, not by the case-level verdict.
        _write_row(ws, row, [relative, digest, recomputed, result],
                   mono_cols=(2, 3), fill=RESULT_FILLS.get(result))
        ws.cell(row=row, column=4).font = _BOLD
        row += 1

    ws.cell(row=row + 1, column=1, value="Verdict").font = _BOLD
    verdict_cell = ws.cell(row=row + 1, column=2, value=integrity["verdict"])
    verdict_cell.font = _BOLD
    verdict_cell.fill = PatternFill("solid", fgColor=VERDICT_FILLS.get(integrity["verdict"], WARN_BG))

    _finish(ws, [52, 68, 68, 18])
    return ws


def _sheet_sources(wb, data):
    ws = wb.create_sheet("Sources")
    _header(ws, ["#", "Loader", "Status", "Verification tier", "Detail"],
            [5, 26, 22, 22, 88])

    earned = {v["loader"]: v["tier"] for v in data["verifications"]}
    statuses = data["statuses"]

    for idx, status in enumerate(statuses, start=1):
        loader = status.get("loader", "")
        tier = status.get("verification_tier") or earned.get(loader, "")
        fill = STATUS_FILLS.get(status.get("status", ""))
        _write_row(ws, idx + 1,
                   [idx, loader, status.get("status", ""), tier, status.get("detail", "")],
                   fill=fill)
        ws.cell(row=idx + 1, column=3).font = _BOLD

    _finish(ws, [5, 26, 22, 22, 88])
    return ws


def _sheet_artifacts(wb, data):
    ws = wb.create_sheet("Artifacts")
    _header(ws, ["Artifact ID", "Loader", "Remote path", "Local path",
                 "Size (bytes)", "SHA-256", "Acquired at", "Status", "Detail"],
            [18, 20, 40, 46, 14, 68, 22, 20, 40])

    artifacts = data["artifacts"]
    for idx, art in enumerate(artifacts, start=1):
        size = art.get("size_bytes")
        _write_row(ws, idx + 1, [
            art.get("artifact_id", ""), art.get("loader", ""),
            art.get("remote_path", ""), art.get("local_path", ""),
            int(size) if isinstance(size, (int, float)) else "",
            art.get("sha256", ""), art.get("acquired_at", ""),
            art.get("status", ""), art.get("detail", ""),
        ], mono_cols=(6,), fill=STATUS_FILLS.get(art.get("status", "")))

    if not artifacts:
        ws.cell(row=2, column=1,
                value="No artifact has been recorded for this case.").font = _BOLD
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)

    _finish(ws, [18, 20, 40, 46, 14, 68, 22, 20, 40])
    return ws


def _sheet_custody(wb, data):
    ws = wb.create_sheet("Custody")
    _header(ws, ["Timestamp (UTC)", "Action", "Actor", "Detail", "Tool", "Version"],
            [24, 22, 22, 74, 14, 10])

    for idx, event in enumerate(data["custody"], start=1):
        _write_row(ws, idx + 1, [
            event.get("timestamp", ""), event.get("action", ""),
            event.get("actor", ""), event.get("detail", ""),
            event.get("tool", ""), event.get("tool_version", ""),
        ])

    _finish(ws, [24, 22, 22, 74, 14, 10])
    return ws


# --------------------------------------------------------------------------
# build
# --------------------------------------------------------------------------

def build(case_id: str, output: Path | None = None) -> Path:
    """Build the verification workbook for ``case_id``. Returns the written path."""
    data = _load(case_id)
    data["artifacts"] = evidence_store.read_artifacts(case_id)

    # Read-only integrity: recomputes hashes without touching the case tree.
    integrity = evidence_store.compute_integrity(case_id)

    wb = Workbook()
    _sheet_summary(wb, data, integrity, case_id)
    _sheet_integrity(wb, data, integrity, case_id)
    _sheet_sources(wb, data)
    _sheet_artifacts(wb, data)
    _sheet_custody(wb, data)

    if output is None:
        out_dir = _ROOT / "reports"
        out_dir.mkdir(parents=True, exist_ok=True)
        output = out_dir / f"{case_id}_verification.xlsx"
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    sidecar = output.with_suffix(output.suffix + ".sha256")
    sidecar.write_text(
        f"{evidence_store.hash_file(output)}  {output.name}\n", encoding="utf-8")
    return output


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Build the XLSX verification workbook.")
    parser.add_argument("case_id", help="case identifier, e.g. CASE-BASELINE-001")
    parser.add_argument("-o", "--output", type=Path, default=None,
                        help="output .xlsx path (default: reports/<case>_verification.xlsx)")
    args = parser.parse_args(argv)

    path = build(args.case_id, args.output)
    print(f"wrote {path}")
    print(f"sidecar {path}.sha256")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())