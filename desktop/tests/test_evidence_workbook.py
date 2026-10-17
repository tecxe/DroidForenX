"""
XLSX verification workbook tests.

The workbook is a report artifact an examiner reads and can re-check by hand, so
these tests pin the properties that make it trustworthy:

  1. The Integrity sheet's "Re-computed SHA-256" column holds a hash actually
     computed from the file on disk - not a copy of the manifest digest. A
     tampered artifact must show a *different* digest, a missing one must show
     nothing at all.
  2. A case with no manifest reports NO_MANIFEST instead of implying verification.
  3. Building the workbook is read-only: the case tree and the custody log are
     byte-for-byte unchanged afterwards.

Run with:  python -m unittest desktop.tests.test_evidence_workbook -v
"""

import hashlib
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from desktop.core import evidence_store
from desktop.report import evidence_workbook


class EvidenceWorkbookTest(unittest.TestCase):
    """Each test gets an isolated evidence root so nothing touches ~/.droidforenx."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="dfx-wb-test-")
        self.evidence_root = Path(self.tmp) / "root"
        self.evidence_root.mkdir()
        self._original_root = evidence_store.EVIDENCE_ROOT
        evidence_store.EVIDENCE_ROOT = self.evidence_root
        self.case_id = "CASE-WB-001"

    def tearDown(self):
        evidence_store.EVIDENCE_ROOT = self._original_root
        evidence_store.set_active_case(None)
        shutil.rmtree(self.tmp, ignore_errors=True)

    # -- helpers ---------------------------------------------------------

    def _make_artifact(self, name, content):
        path = evidence_store.acquired_dir(self.case_id) / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        return path

    def _build(self):
        out = Path(self.tmp) / f"{self.case_id}_verification.xlsx"
        evidence_workbook.build(self.case_id, output=out)
        return load_workbook(out)

    def _integrity_rows(self, wb):
        ws = wb["Integrity"]
        rows = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[3] in ("OK", "MISMATCH", "MISSING"):
                rows[Path(row[0]).name] = row
        return rows

    # -- 1. the re-computed column is a real re-hash ----------------------

    def test_recomputed_column_shows_on_disk_hash(self):
        evidence_store.open_case(self.case_id, operator="test")
        ok = self._make_artifact("ok.db", b"SQLite format 3\x00clean")
        bad = self._make_artifact("bad.db", b"SQLite format 3\x00original")
        gone = self._make_artifact("gone.db", b"SQLite format 3\x00deleted")
        evidence_store.record_acquired(self.case_id, "/data/ok.db", ok, loader="whatsapp")
        evidence_store.record_acquired(self.case_id, "/data/bad.db", bad, loader="telegram")
        evidence_store.record_acquired(self.case_id, "/data/gone.db", gone, loader="instagram")

        bad.write_bytes(b"SQLite format 3\x00TAMPERED")
        gone.unlink()

        rows = self._integrity_rows(self._build())

        path, manifest, recomputed, result = rows["ok.db"]
        self.assertEqual(result, "OK")
        self.assertEqual(recomputed, manifest,
                         "an intact artifact must re-hash to its manifest digest")
        self.assertEqual(recomputed, evidence_store.hash_file(ok))

        path, manifest, recomputed, result = rows["bad.db"]
        self.assertEqual(result, "MISMATCH")
        self.assertNotEqual(recomputed, manifest,
                            "a tampered artifact must not echo the manifest digest")
        self.assertEqual(recomputed, evidence_store.hash_file(bad),
                         "the re-computed column must be the real on-disk hash")

        path, manifest, recomputed, result = rows["gone.db"]
        self.assertEqual(result, "MISSING")
        self.assertFalse(recomputed,
                         "a missing artifact must leave the re-hash blank, not copy the manifest")

    def test_compute_integrity_exposes_actual_hashes(self):
        evidence_store.open_case(self.case_id)
        a = self._make_artifact("a.db", b"alpha")
        b = self._make_artifact("b.db", b"bravo")
        evidence_store.record_acquired(self.case_id, "/a", a, loader="x")
        evidence_store.record_acquired(self.case_id, "/b", b, loader="x")
        b.unlink()

        result = evidence_store.compute_integrity(self.case_id)
        actual = result["actual"]
        self.assertEqual(actual["acquired/a.db"], evidence_store.hash_file(a))
        self.assertIsNone(actual["acquired/b.db"],
                          "a missing artifact must map to None, not a stale digest")

    # -- 2. no manifest is stated, not implied ---------------------------

    def test_no_manifest_verdict_is_reported(self):
        evidence_store.open_case(self.case_id)
        wb = self._build()

        integrity = wb["Integrity"]
        note = integrity.cell(row=2, column=1).value
        self.assertIn("NO_MANIFEST", note)

        summary = wb["Case Summary"]
        verdict_cell = None
        for r in range(1, summary.max_row + 1):
            if summary.cell(row=r, column=1).value == "INTEGRITY":
                verdict_cell = summary.cell(row=r, column=2)
        self.assertIsNotNone(verdict_cell, "the summary must carry an INTEGRITY row")
        self.assertEqual(verdict_cell.value, "NO_MANIFEST")

    # -- 3. building a workbook never alters evidence --------------------

    def test_build_does_not_touch_case_tree(self):
        evidence_store.open_case(self.case_id, operator="test")
        artifact = self._make_artifact("evidence.db", b"original content")
        evidence_store.record_acquired(self.case_id, "/data/x", artifact, loader="test")

        case_dir = evidence_store.case_dir(self.case_id)

        def snapshot():
            state = {}
            for path in sorted(case_dir.rglob("*")):
                if path.is_file():
                    state[str(path.relative_to(case_dir))] = hashlib.sha256(
                        path.read_bytes()).hexdigest()
            return state

        before = snapshot()
        custody_before = len(evidence_store.read_custody(self.case_id))

        self._build()

        self.assertEqual(snapshot(), before,
                         "building a workbook must not modify, add or remove any case file")
        self.assertEqual(len(evidence_store.read_custody(self.case_id)), custody_before,
                         "building a workbook must not append a custody event")

    def test_build_writes_sidecar_hash(self):
        evidence_store.open_case(self.case_id)
        artifact = self._make_artifact("evidence.db", b"content")
        evidence_store.record_acquired(self.case_id, "/data/x", artifact, loader="test")

        out = Path(self.tmp) / "case.xlsx"
        evidence_workbook.build(self.case_id, output=out)

        sidecar = out.with_suffix(".xlsx.sha256")
        self.assertTrue(sidecar.exists())
        digest, _, name = sidecar.read_text(encoding="utf-8").strip().partition("  ")
        self.assertEqual(name, out.name)
        self.assertEqual(digest, evidence_store.hash_file(out))


if __name__ == "__main__":
    unittest.main()
